"""
CV Generator — produces a single-page .docx from structured JSON.
Uses the same clean format: Calibri, section dividers, bold labels.
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_border(cell, **kwargs):
    """Set cell border (unused but kept for future)."""
    pass


def add_bottom_border(paragraph):
    """Add a thin bottom border to a paragraph (section divider)."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_tab_stop_right(paragraph, position_inches=6.5):
    """Add a right-aligned tab stop."""
    pPr = paragraph._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(int(position_inches * 1440)))
    tabs.append(tab)
    pPr.append(tabs)


def set_paragraph_spacing(paragraph, before=0, after=0, line=None):
    """Set paragraph spacing in points."""
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = Pt(line)


def add_run(paragraph, text, bold=False, italic=False, size=10, font="Calibri"):
    """Add a formatted run to a paragraph."""
    run = paragraph.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return run


def generate_cv_docx(cv_data, output_path):
    """Generate a .docx CV from structured data."""
    doc = Document()

    # Page setup
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.375)
    section.bottom_margin = Inches(0.25)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    # Set default font
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(10)

    tab_pos = 7.0  # right tab position in inches

    # ── NAME ──
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=2)
    add_run(p, cv_data["name"], bold=True, size=14)

    # ── CONTACT ──
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=4)
    add_run(p, cv_data["contact"], size=9)

    # ── PROFESSIONAL EXPERIENCE ──
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=7, after=2)
    add_bottom_border(p)
    add_run(p, "PROFESSIONAL EXPERIENCE", bold=True, size=11)

    for exp in cv_data.get("experience", []):
        # Company header
        company_text = exp["company"]
        if exp.get("company_descriptor"):
            company_text += f" {exp['company_descriptor']}"

        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=5, after=1)
        add_tab_stop_right(p, tab_pos)
        add_run(p, company_text, bold=True, size=10)
        add_run(p, "\t", size=10)
        add_run(p, exp["location"], italic=True, size=10)

        # Role titles
        for role in exp.get("roles", []):
            p = doc.add_paragraph()
            set_paragraph_spacing(p, before=0.5, after=1)
            add_tab_stop_right(p, tab_pos)
            add_run(p, role["title"], italic=True, size=10)
            add_run(p, "\t", size=10)
            add_run(p, role["dates"], size=10)

        # Bullets
        for bullet in exp.get("bullets", []):
            p = doc.add_paragraph()
            set_paragraph_spacing(p, before=0.5, after=0.5)
            pf = p.paragraph_format
            pf.left_indent = Inches(0.25)
            pf.first_line_indent = Inches(-0.125)

            add_run(p, "\u2022 ", size=10)
            add_run(p, f"{bullet['label']}: ", bold=True, size=10)
            add_run(p, bullet["text"], size=10)

    # ── EDUCATION ──
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=7, after=2)
    add_bottom_border(p)
    add_run(p, "EDUCATION", bold=True, size=11)

    for edu in cv_data.get("education", []):
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=5, after=0.5)
        add_tab_stop_right(p, tab_pos)
        add_run(p, edu["institution"], bold=True, size=10)
        add_run(p, "\t", size=10)
        add_run(p, edu["location"], italic=True, size=10)

        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=0.5, after=0.5)
        add_tab_stop_right(p, tab_pos)
        add_run(p, edu["degree"], size=10)
        add_run(p, "\t", size=10)
        add_run(p, edu["dates"], size=10)

    # ── SKILLS ──
    skills = cv_data.get("skills", {})
    if skills:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=7, after=2)
        add_bottom_border(p)
        add_run(p, "SKILLS & EXPERTISE", bold=True, size=11)

        for key, label in [
            ("languages", "Languages"),
            ("technical", "Technical"),
            ("domain", "Domain"),
        ]:
            val = skills.get(key)
            if val:
                p = doc.add_paragraph()
                set_paragraph_spacing(p, before=1, after=0.5)
                add_run(p, f"{label}: ", bold=True, size=10)
                add_run(p, val, size=10)

    doc.save(output_path)
    return output_path


def estimate_fits_one_page(cv_data):
    """
    Estimate whether the CV fits on one page by calculating total content height.
    Uses the exact same spacing, font sizes, and margins as generate_cv_docx.

    Returns: (fits: bool, used_pts: float, available_pts: float)
    """
    # Page geometry (matching generate_cv_docx exactly)
    AVAILABLE_HEIGHT_PTS = (11 - 0.375 - 0.25) * 72  # 747 points
    BODY_WIDTH_PTS = (8.5 - 0.75 - 0.75) * 72        # 504 points
    BULLET_WIDTH_PTS = (7.0 - 0.25) * 72              # 486 points (after indent)

    # Calibri average char width: ~0.5 × font_size (conservative — slightly wide)
    CHAR_WIDTH_FACTOR = 0.50

    def line_height(font_size_pt):
        """Single-spaced line height for Calibri (Word default ~1.15x)."""
        return font_size_pt * 1.15

    def text_lines(text, font_size_pt, width_pts):
        """How many lines does this text occupy at the given width?"""
        if not text:
            return 1
        char_w = font_size_pt * CHAR_WIDTH_FACTOR
        chars_per_line = max(1, int(width_pts / char_w))
        return max(1, -(-len(text) // chars_per_line))  # ceil division

    def para_height(text, font_size_pt, width_pts, space_before=0, space_after=0):
        """Total height of one paragraph in points."""
        n_lines = text_lines(text, font_size_pt, width_pts)
        return space_before + (n_lines * line_height(font_size_pt)) + space_after

    h = 0.0

    # Name (14pt, before=0, after=2)
    h += para_height(cv_data.get("name", ""), 14, BODY_WIDTH_PTS, 0, 2)

    # Contact (9pt, before=0, after=4)
    h += para_height(cv_data.get("contact", ""), 9, BODY_WIDTH_PTS, 0, 4)

    # "PROFESSIONAL EXPERIENCE" section header (11pt, before=7, after=2)
    h += para_height("PROFESSIONAL EXPERIENCE", 11, BODY_WIDTH_PTS, 7, 2)

    for exp in cv_data.get("experience", []):
        # Company header — tab-separated, always 1 line (10pt, before=5, after=1)
        h += 5 + line_height(10) + 1

        # Role titles — tab-separated, always 1 line each (10pt, before=0.5, after=1)
        for _ in exp.get("roles", []):
            h += 0.5 + line_height(10) + 1

        # Bullets (10pt, before=0.5, after=0.5, indented width)
        for bullet in exp.get("bullets", []):
            full_text = f"\u2022 {bullet.get('label', '')}: {bullet.get('text', '')}"
            h += para_height(full_text, 10, BULLET_WIDTH_PTS, 0.5, 0.5)

    # "EDUCATION" section header (11pt, before=7, after=2)
    h += para_height("EDUCATION", 11, BODY_WIDTH_PTS, 7, 2)

    for edu in cv_data.get("education", []):
        # Institution — tab-separated, 1 line (10pt, before=5, after=0.5)
        h += 5 + line_height(10) + 0.5
        # Degree — tab-separated, 1 line (10pt, before=0.5, after=0.5)
        h += 0.5 + line_height(10) + 0.5

    # "SKILLS & EXPERTISE" section header (11pt, before=7, after=2)
    skills = cv_data.get("skills", {})
    if skills:
        h += para_height("SKILLS & EXPERTISE", 11, BODY_WIDTH_PTS, 7, 2)

        for key in ["languages", "technical", "domain"]:
            val = skills.get(key)
            if val:
                full_text = f"{key}: {val}"
                h += para_height(full_text, 10, BODY_WIDTH_PTS, 1, 0.5)

    fits = h <= AVAILABLE_HEIGHT_PTS
    return fits, round(h, 1), round(AVAILABLE_HEIGHT_PTS, 1)
